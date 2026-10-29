"""Reconstruct genuine TP1-to-TP2 redlines from the supplied team norms; retain real Word revisions."""
from pathlib import Path
from datetime import datetime, timezone
import difflib
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
# Text transcribed from the team's submitted TP1 norms. Mapping preserves the TP2 team's order.
OLD = [
 'Hold 2 “standup” meetings a week. (Required by course, provides clarity between team’s expectations)',
 'Being responsive; WITHIN 6 hours. (Ensures quick communication for effectively mitigating problems)',
 'Being punctual for group meetings, and participating in the discussions. (To make sure we can coordinate efficiently)',
 'Being respectful and understanding. (To avoid miscommunication and have better team spirit)',
 'Attempt to reach a compromise for disagreements. (Human civilization is built, and runs on this)',
 'Complete assigned work by the team-agreed internal deadline. (Provides sufficient time for review, testing, integration, and correction before the official submission deadline)',
 'Coordinate changes that affect shared code or another member’s work before implementing them. (Prevents conflicting implementations and reduces integration problems)',
 'Review and test integrated work as a team before submission. (Ensures the final application and deliverables function together correctly rather than merely having individually completed components)',
]
# Current TP2 wording and explanations come from the supplied TP2 document, not invented meeting history.
NORMS = [
 (0, 'Hold at least two 15-minute standup meetings each week, at least 24 hours apart, and record each one.',
  'Required by the course for at least six standups per phase; regular check-ins surface problems early.'),
 (7, 'Review and test the integrated work together before every submission.',
  'Confirms the combined application and deliverables meet the team’s expectations, not just each person’s piece. Moved from TP1 norm 8 to emphasize integration.'),
 (1, 'Respond to team messages within 6 hours during the day, and by the next morning for messages sent late at night.',
  'Quick replies let the team resolve problems before they delay others. The overnight exception makes the response expectation clear.'),
 (2, 'Arrive on time to team meetings and take part in the discussion.',
  'Lets the team coordinate efficiently and make decisions with everyone’s input.'),
 (3, 'Be respectful and understanding toward every team member.',
  'Reduces miscommunication and builds a stronger working relationship.'),
 (4, 'Work toward a compromise when members disagree.',
  'Keeps disagreements from stalling progress and gives everyone’s view a fair hearing.'),
 (5, 'Complete assigned work by the team’s internal deadline.',
  'Leaves enough time for review, testing, integration, and corrections before the official deadline.'),
 (6, 'Coordinate before changing shared code or another member’s work.',
  'Prevents conflicting changes and helps the separately built pieces work together correctly.'),
 (None, 'Tell the team as soon as you are blocked or may miss an internal deadline.',
  'Lets others adjust or help in time, instead of finding out at the deadline.'),
 (None, 'Pull the latest code before starting work, and commit only source files, not compiled bin files.',
  'Avoids merge conflicts and keeps the shared code consistent for everyone.'),
]


def tokens(text):
    return re.findall(r'\S+\s*', text)


def changes(old, new):
    a, b = tokens(old), tokens(new)
    for op, i, j, k, l in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if op == 'equal':
            yield 'equal', ''.join(a[i:j])
        else:
            if op in ('replace', 'delete'):
                yield 'delete', ''.join(a[i:j])
            if op in ('replace', 'insert'):
                yield 'insert', ''.join(b[k:l])


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.7)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for style_name in ('Heading 1', 'Heading 2', 'Heading 3'):
        doc.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    title = doc.add_paragraph('TP2 Team Norms', 'Title')
    for run in title.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.size = Pt(21)
    border = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'nil')
    border.append(bottom)
    title._p.get_or_add_pPr().append(border)
    doc.add_paragraph('Team 53 | CSE 360 | Updated October 28, 2026')
    doc.add_paragraph('These norms update the team’s submitted TP1 norms using the supplied TP2 wording. '
                      'Deleted TP1 wording is struck through; inserted TP2 wording is '
                      'underlined. Each norm includes the reason for the change. The editable source '
                      'retains actual Word tracked insertions and deletions.')
    settings = doc.settings.element
    for tag in ('trackRevisions', 'revisionView'):
        node = OxmlElement('w:' + tag)
        if tag == 'revisionView':
            for key in ('markup', 'insDel', 'formatting'):
                node.set(qn('w:' + key), 'true')
        settings.append(node)
    revision_id = 0
    when = datetime.now(timezone.utc).isoformat(timespec='seconds').replace('+00:00', 'Z')
    for number, (previous, norm, reason) in enumerate(NORMS, 1):
        doc.add_paragraph(f'{number}. ' + ('New norm' if previous is None else f'Updates TP1 norm {previous + 1}'), 'Heading 2')
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        for kind, text in changes(OLD[previous] if previous is not None else '', norm):
            if kind == 'equal':
                paragraph.add_run(text)
                continue
            revision_id += 1
            revision = OxmlElement('w:del' if kind == 'delete' else 'w:ins')
            revision.set(qn('w:id'), str(revision_id))
            revision.set(qn('w:author'), 'TP2 integration review')
            revision.set(qn('w:date'), when)
            run = OxmlElement('w:r')
            properties = OxmlElement('w:rPr')
            color = OxmlElement('w:color')
            color.set(qn('w:val'), 'B42318' if kind == 'delete' else '1454A3')
            properties.append(color)
            decoration = OxmlElement('w:strike' if kind == 'delete' else 'w:u')
            decoration.set(qn('w:val'), 'true' if kind == 'delete' else 'single')
            properties.append(decoration)
            run.append(properties)
            node = OxmlElement('w:delText' if kind == 'delete' else 'w:t')
            node.set(qn('xml:space'), 'preserve')
            node.text = text
            run.append(node)
            revision.append(run)
            paragraph._p.append(revision)
        p = doc.add_paragraph()
        p.add_run('Reason for change: ').bold = True
        p.add_run(reason)
    doc.add_paragraph('Team members: Vishwam Gupta; Joshua Luther; Kanish Garg; '
                      'Aditya Gupta (agupt515); Aditya Gupta (agupt545).')
    doc.core_properties.title = 'TP2 Team Norms'
    doc.core_properties.author = 'Team 53'
    doc.core_properties.comments = 'Redlines reconstructed from the actual supplied TP1 and TP2 documents on October 28, 2026.'
    target = ROOT / 'TP2 Team Norms.docx'
    doc.save(target)
    print('Wrote', target, 'with', revision_id, 'actual tracked revisions')


if __name__ == '__main__':
    build()
